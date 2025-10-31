import re
from pathlib import Path
from typing import Iterable

import requests
from docx import Document
from bs4 import BeautifulSoup
from bs4.element import Tag

from database import is_chapter_in_db, add_chapter_in_db

SKIP_LINES = {"Предыдущая глава", "Следующая глава"}

baseDir = Path(__file__).resolve().parents[1]
chapterDir = baseDir / "data" / "chapters"
if not chapterDir.exists():
    chapterDir.mkdir(parents=True, exist_ok=True)

def write_to_txt_file(txt_path: Path, chapter_title: str, chapter_text: Iterable[Tag]) -> None:
    if not chapter_text:
        return

    with txt_path.open("w", encoding="utf-8") as file:
        file.write(f'{chapter_title.strip()}\n\n')
        
        for text_line in chapter_text:
            filtered_line = text_line.get_text(strip=True)
            if not filtered_line or filtered_line in SKIP_LINES:
                continue
            file.write(f'{filtered_line}\n')


def write_to_docx_file(docx_path: Path, chapter_title: str, chapter_text: Iterable[Tag]) -> None:
    if not chapter_text:
        return
    
    document = Document()
    heading = document.add_heading(level=1)
    run = heading.add_run(chapter_title.strip())
    run.bold = True
    
    for text_line in chapter_text:
        filtered_line = text_line.get_text(strip=True)
        if not filtered_line or filtered_line in SKIP_LINES:
            continue
        document.add_paragraph(filtered_line, style="No Spacing")
    document.save(str(docx_path))


def parsChapter(url: str) -> bool:
    headers = {
        "Accept": "*/*",
        "User-Agent": "Mozilla/5.0 (X11; Ubuntu; Linux x86_64; rv:137.0) Gecko/20100101 Firefox/137.0"
    }
    try:
        req = requests.get(url, headers=headers)
        src = req.text
        soup = BeautifulSoup(src, "lxml")

        content = soup.find(class_="tl_article_content")
        if content is None:
            raise ValueError("Missing .tl_article_content section in page")

        chapter_title = content.find("h1").text
        chapter_text = content.find_all("p")
        chapter_number =  re.sub(':.*', '', chapter_title)[6:]

        if is_chapter_in_db(chapter_number): 
            print(f'Chapter "{chapter_title}" skipped: already exists!')
            return True

        path = chapterDir / chapter_number
        if not path.exists():
            path.mkdir(parents=True, exist_ok=True)

        txt_path = path / f"{chapter_number}.txt"
        docx_path = path / f"{chapter_number}.docx"

        write_to_txt_file(txt_path, chapter_title, chapter_text)
        write_to_docx_file(docx_path, chapter_title, chapter_text)
        

        add_chapter_in_db(chapter_number, chapter_title, url, str(path))
        return True
    except Exception as e:
        print(f"Error parsing {url}: {e}")
        return False